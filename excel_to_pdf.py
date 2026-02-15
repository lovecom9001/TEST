"""
엑셀 파일 자동화 스크립트
- 엑셀 파일 열기
- 특정 시트 불러오기
- 특정 셀에 데이터 입력
- 특정 시트를 PDF로 저장
"""

import openpyxl
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle


def open_workbook(file_path: str) -> openpyxl.Workbook:
    """엑셀 파일을 열고 Workbook 객체를 반환한다."""
    wb = openpyxl.load_workbook(file_path)
    print(f"엑셀 파일 열기 완료: {file_path}")
    print(f"시트 목록: {wb.sheetnames}")
    return wb


def get_sheet(wb: openpyxl.Workbook, sheet_name: str):
    """특정 시트를 불러온다."""
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"시트 '{sheet_name}'이(가) 존재하지 않습니다. 사용 가능한 시트: {wb.sheetnames}")
    ws = wb[sheet_name]
    print(f"시트 불러오기 완료: {sheet_name}")
    return ws


def write_cell(ws, cell: str, value):
    """특정 셀에 데이터를 입력한다. cell은 'A1', 'B2' 같은 형식."""
    ws[cell] = value
    print(f"셀 {cell}에 데이터 입력 완료: {value}")


def write_cells(ws, data: dict):
    """여러 셀에 데이터를 한번에 입력한다. data = {'A1': 값, 'B2': 값, ...}"""
    for cell, value in data.items():
        write_cell(ws, cell, value)


def save_workbook(wb: openpyxl.Workbook, file_path: str):
    """변경된 엑셀 파일을 저장한다."""
    wb.save(file_path)
    print(f"엑셀 파일 저장 완료: {file_path}")


def sheet_to_pdf(ws, pdf_path: str):
    """시트 내용을 PDF로 저장한다.

    Args:
        ws: openpyxl 워크시트 객체
        pdf_path: 저장할 PDF 파일 경로
    """
    max_col = ws.max_column
    max_row = ws.max_row

    if max_row is None or max_col is None or max_row == 0:
        print("시트에 데이터가 없습니다.")
        return

    # 시트 데이터를 2차원 리스트로 변환
    table_data = []
    for row_idx in range(1, max_row + 1):
        row = []
        for col_idx in range(1, max_col + 1):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            row.append(str(cell_value) if cell_value is not None else "")
        table_data.append(row)

    # PDF 문서 생성 (가로 방향)
    doc = SimpleDocTemplate(pdf_path, pagesize=landscape(A4))

    # 테이블 생성
    table = Table(table_data)

    # 테이블 스타일 설정
    style = TableStyle([
        # 헤더 행 스타일
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTSIZE", (0, 0), (-1, 0), 12),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        # 데이터 행 스타일
        ("FONTSIZE", (0, 1), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 6),
        ("TOPPADDING", (0, 1), (-1, -1), 6),
        # 전체 테이블 스타일
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
        # 짝수 행 배경색
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.Color(0.93, 0.93, 0.93)]),
    ])
    table.setStyle(style)

    # PDF 빌드
    doc.build([table])
    print(f"PDF 저장 완료: {pdf_path}")


# ──────────────────────────────────────────────
# 사용 예시 (직접 실행 시)
# ──────────────────────────────────────────────
if __name__ == "__main__":
    import os

    # ── 설정값 (필요에 따라 수정) ──────────────
    EXCEL_FILE = "sample.xlsx"       # 엑셀 파일 경로
    SHEET_NAME = "Sheet1"            # 작업할 시트 이름
    PDF_OUTPUT = "output.pdf"        # PDF 저장 경로

    # 입력할 데이터 (셀 주소: 값)
    CELL_DATA = {
        "A1": "Name",
        "B1": "Age",
        "C1": "Email",
        "A2": "Hong Gildong",
        "B2": 30,
        "C2": "hong@example.com",
        "A3": "Kim Cheolsu",
        "B3": 25,
        "C3": "kim@example.com",
    }
    # ──────────────────────────────────────────

    # 파일이 없으면 새로 생성
    if not os.path.exists(EXCEL_FILE):
        print(f"'{EXCEL_FILE}' 파일이 없어 새로 생성합니다.")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        wb.save(EXCEL_FILE)

    # 1. 엑셀 파일 열기
    wb = open_workbook(EXCEL_FILE)

    # 2. 특정 시트 불러오기
    ws = get_sheet(wb, SHEET_NAME)

    # 3. 특정 셀에 데이터 입력
    write_cells(ws, CELL_DATA)

    # 4. 엑셀 파일 저장
    save_workbook(wb, EXCEL_FILE)

    # 5. 시트를 PDF로 저장
    sheet_to_pdf(ws, PDF_OUTPUT)

    print("\n모든 작업이 완료되었습니다!")
